import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

rules = []

def add(domain, section, subcategory, condition, keywords='', notes=''):
    rules.append({
        'rule_id': f'R{len(rules)+1:03d}',
        'domain': domain,
        'section': section,
        'subcategory': subcategory,
        'mor_condition': condition,
        'suggested_label': 'MOR',
        'suggested_category': domain,
        'suggested_subcategory': subcategory,
        'keywords_or_triggers': keywords,
        'notes_or_exceptions': notes,
        'source_document': 'Rules for Reportable Occurences.pdf'
    })

# I. Aircraft Technical - Structural
for c,k in [
('Damage to a principal structural element that has not been qualified as damage tolerant / life limited element.', 'principal structural element, damage tolerant, life limited'),
('Damage or defect exceeding allowed tolerances to a structural element whose failure could reduce structural stiffness such that required aeroelastic characteristics are no longer achieved.', 'structural defect, tolerance, stiffness, aeroelastic'),
('Damage to or defect of a structural element which could result in liberation of items of mass that may injure aircraft occupants.', 'structural element, liberation, loose part, injury'),
('Damage to or defect of a structural element which could jeopardise proper operation of systems.', 'structural damage, system operation'),
('Loss of any part of the aircraft structure in flight.', 'loss of structure, in flight, aircraft part detached')]:
    add('Aircraft Technical','I-A','Structural',c,k)

# Systems generic
for c,k in [
('Loss, significant malfunction or defect of any system, sub-system or equipment when standard operating procedures could not be satisfactorily accomplished.', 'system loss, malfunction, defect, SOP not accomplished'),
('Inability of crew to control a system, including uncommanded actions or incorrect/incomplete response.', 'unable to control, uncommanded, incomplete response'),
('Failure or malfunction of protection device or emergency system associated with a system.', 'protection device, emergency system failure'),
('Loss of redundancy of a system.', 'loss of redundancy'),
('Operation of any primary warning system associated with aircraft systems/equipment unless conclusively false and not hazardous.', 'primary warning, aircraft system warning'),
('Leakage of hydraulic fluid, fuel, oil or other fluids that may result in fire hazard, hazardous contamination or risk to occupants.', 'leakage, hydraulic, fuel, oil, fluid, fire hazard, contamination'),
('Malfunction or defect of any indication system resulting in possibility of misleading indications to crew.', 'indication system, misleading indication'),
('Any failure, malfunction or defect in a system during critical phase of flight.', 'failure, malfunction, critical phase of flight'),
('Flight controls malfunction.', 'flight control malfunction')]:
    add('Aircraft Technical','I-B','Systems - Generic',c,k)

systems = {
'Air conditioning/ventilation': [('Complete loss of avionics cooling.','avionics cooling loss'),('Depressurization.','depressurization')],
'Auto-flight system': [('Failure of auto-flight system to achieve intended operation while engaged.','autoflight failure, autopilot'),('Difficulty controlling aircraft linked to auto-flight system functioning.','autoflight, difficulty control'),('Failure of any auto-flight system disconnect device.','autoflight disconnect failure'),('Uncommanded auto-flight mode change.','uncommanded mode change')],
'Communications': [('Failure or defect of passenger address system resulting in loss or inaudible passenger address.','passenger address, PA failure'),('Total loss of communication in flight.','communication loss, radio loss')],
'Electrical system': [('Loss of one electrical distribution system AC/DC.','AC DC distribution loss'),('Total loss or loss of more than one electrical generation system.','electrical generation loss'),('Failure of backup/emergency electrical generating system.','emergency generator failure')],
'Cockpit/Cabin/Cargo': [('Pilot seat control loss during flight.','pilot seat control loss'),('Failure of emergency system/equipment including evacuation signalling, exit doors or emergency lighting.','emergency equipment, exit door, emergency lighting'),('Loss of retention capability of cargo loading system.','cargo loading retention loss')],
'Fire protection system': [('Fire warnings, except those immediately confirmed as false.','fire warning'),('Undetected failure/defect of fire or smoke detection/protection system that could reduce detection/protection.','fire detection, smoke detection, protection failure'),('Absence of warning in case of actual fire or smoke.','actual fire, smoke, no warning')],
'Flight controls': [('Asymmetry of flaps, slats, spoilers etc.','flap asymmetry, slat asymmetry, spoiler asymmetry'),('Limitation of movement, stiffness or poor/delayed response in primary flight control systems or sub-systems.','flight control stiffness, delayed response'),('Flight control surface runaway.','control surface runaway'),('Flight control surface vibration.','control surface vibration'),('Mechanical flight control disconnection or failure.','mechanical flight control failure'),('Significant interference with normal control or degradation of flying qualities.','degraded flying qualities')],
'Fuel system': [('Fuel quantity indicating system malfunction causing total loss or erroneous indicated fuel quantity on board.','fuel quantity indication, erroneous fuel'),('Fuel leakage causing loss, fire hazard or significant contamination.','fuel leak, fire hazard'),('Fuel jettisoning system malfunction/defect causing inadvertent significant loss, fire hazard, hazardous contamination or inability to jettison fuel.','fuel jettison, inadvertent loss'),('Fuel system malfunction/defect significantly affecting fuel supply and/or distribution.','fuel supply, fuel distribution'),('Inability to transfer or use total quantity of usable fuel.','unable transfer fuel, unusable fuel')],
'Hydraulics': [('Loss of hydraulic system.','hydraulic system loss'),('Leakage of hydraulic fluid.','hydraulic leak'),('Loss of more than one hydraulic circuit.','hydraulic circuit loss'),('Failure of backup hydraulic system.','backup hydraulic failure'),('Inadvertent Ram Air Turbine extension.','RAT extension, ram air turbine')],
'Ice detection/protection system': [('Undetected loss or reduced performance of anti-ice/de-ice system.','anti ice, de ice, reduced performance'),('Loss of more than one probe heating system.','probe heat loss'),('Inability to obtain symmetrical wing de-icing.','asymmetric deicing'),('Abnormal ice accumulation significantly affecting performance or handling.','ice accumulation, handling'),('Crew vision significantly affected.','crew vision, ice')],
'Indicating/warning/recording systems': [('Malfunction/defect of any indicating system with possibility of misleading indications to crew.','misleading indication'),('Loss or malfunction of more than one display unit or computer display/warning function in glass cockpit.','display unit loss, glass cockpit')],
'Landing gear/brakes/tyres': [('Brake fire.','brake fire'),('Significant loss of braking action.','braking action loss'),('Unsymmetrical braking.','asymmetric braking'),('Failure of landing gear free fall extension system.','free fall extension failure'),('Unwanted gear or gear doors extension/retraction.','unwanted gear extension, gear door'),('Tyre burst.','tyre burst, tire burst')],
'Navigation systems': [('Total loss or multiple navigation equipment failures.','navigation failure'),('Total loss or multiple air data system equipment failures.','air data failure'),('Significant misleading navigation indication.','misleading navigation indication'),('Significant navigation errors attributed to incorrect data.','navigation error, incorrect data'),('Unexpected lateral or vertical path deviations not caused by pilot input.','unexpected deviation, lateral path, vertical path')],
'Oxygen': [('For pressurized aircraft: loss of oxygen supply in cockpit.','cockpit oxygen loss'),('Loss of oxygen supply to significant number of passengers, more than 10%.','passenger oxygen loss')],
'Bleed air system': [('Hot bleed air leak causing fire warning or structural damage.','hot bleed air leak'),('Loss of all bleed air systems.','bleed air loss'),('Failure of bleed air leak detection system.','bleed leak detection failure')],
}
for sub, items in systems.items():
    for c,k in items: add('Aircraft Technical','I-B',sub,c,k)

propulsion = [
('Flameout, shutdown or malfunction of any engine.','flameout, engine shutdown, engine malfunction'),
('Overspeed or inability to control speed of any high-speed rotating component.','overspeed, rotating component'),
('Engine/powerplant failure causing non-containment of components/debris.','non-containment, debris'),
('Engine/powerplant failure causing uncontrolled internal or external fire.','engine fire, uncontrolled fire'),
('Engine/powerplant failure causing thrust in a direction different from pilot demand.','wrong thrust direction'),
('Thrust reversing system failing to operate or operating inadvertently.','thrust reverser failure, inadvertent reverse'),
('Inability to control power, thrust or rpm.','unable control thrust, rpm'),
('Failure of engine mount structure.','engine mount failure'),
('Partial or complete loss of a major part of powerplant.','powerplant part loss'),
('Dense visible fumes or toxic product concentration sufficient to incapacitate crew/passengers.','dense fumes, toxic fumes'),
('Inability to shut down an engine using normal procedures.','unable shutdown engine'),
('Inability to restart a serviceable engine.','unable restart engine'),
('Uncommanded thrust/power loss or change classified as loss of thrust/power control.','uncommanded thrust loss'),
('Defects of common origin resulting in in-flight engine shutdown.','in flight engine shutdown'),
('Engine limiter/control device failing when required or operating inadvertently.','engine limiter failure'),
('Exceedance of engine parameters.','engine parameter exceedance'),
('FOD resulting in damage.','FOD damage, foreign object debris'),
('Propeller/powerplant failure causing overspeed of propeller.','propeller overspeed'),
('Propeller/powerplant failure causing excessive drag.','propeller excessive drag'),
('Propeller/powerplant failure causing thrust opposite to pilot command.','opposite thrust'),
('Release of propeller or major portion of propeller.','propeller release'),
('Propeller failure resulting in excessive imbalance.','propeller imbalance'),
('Unintended movement of propeller blades below minimum in-flight low-pitch position.','propeller low pitch'),
('Inability to feather propeller.','unable feather propeller'),
('Inability to command change in propeller pitch.','propeller pitch change failure'),
('Uncommanded change in propeller pitch.','uncommanded pitch'),
('Uncontrollable torque or speed fluctuation.','torque fluctuation, speed fluctuation'),
('Damage/defect of main rotor gearbox/attachment causing possible in-flight rotor separation or rotor control malfunction.','main rotor gearbox, rotor separation'),
('Damage to tail rotor, transmission or equivalent systems.','tail rotor, transmission damage'),
('APU shutdown or failure when APU is required by operational requirements.','APU failure, APU shutdown'),
('Inability to shut down APU.','unable shutdown APU'),
('APU overspeed.','APU overspeed'),
('Inability to start APU when needed for operational reasons.','unable start APU')]
for c,k in propulsion: add('Aircraft Technical','I-B','Propulsion system',c,k)
add('Aircraft Technical','I-C','Human Factors','Any incident where aircraft design feature or inadequacy contributes to hazardous or catastrophic effect.','design inadequacy, hazardous, catastrophic')

# II Flight Operations
ops = [
('Risk of collision with aircraft, terrain or other object, or unsafe situation where avoidance action would have been appropriate.','risk of collision, terrain, object'),('Avoidance manoeuvre required to avoid collision with aircraft, terrain or other object.','avoidance manoeuvre'),('Avoidance manoeuvre to avoid other unsafe situations.','avoidance manoeuvre, unsafe'),('Take-off or landing incidents including precautionary/forced landings; undershoot, overrun, runway excursion, closed/occupied/incorrect runway, runway incursions.','takeoff incident, landing incident, rejected takeoff, runway incursion, overrun, excursion'),('Inability to achieve predicted performance during take-off or initial climb.','performance takeoff, initial climb'),('Critically low fuel quantity or inability to transfer/use total usable fuel.','low fuel, fuel transfer'),('Loss of control, partial or temporary, from any cause.','loss of control'),('Occurrences close to or above V1 causing hazardous/potentially hazardous situation such as rejected take-off, tail strike, engine power loss.','V1, RTO, rejected takeoff, tail strike'),('Unintentional significant deviation from airspeed, intended track or altitude.','airspeed deviation, track deviation, altitude deviation'),('Descent below decision height/altitude or minimum descent height/altitude without required visual reference.','below DH, below MDA, visual reference'),('Loss of position awareness relative to actual position or other aircraft.','position awareness'),('Breakdown in communication between flight crew or with cabin crew, ATC or engineering.','communication breakdown'),('Abnormal runway contact of aircraft.','abnormal runway contact, hard landing'),('Balked landing.','balked landing'),('Exceedance of fuel imbalance limits.','fuel imbalance'),('Incorrect receipt or interpretation of radiotelephony messages.','RT message incorrect'),('Fuel system malfunction/defect affecting fuel supply/distribution.','fuel supply, distribution'),('Aircraft unintentionally departing paved surface.','depart paved surface'),('Collision between aircraft and any aircraft, vehicle or ground object.','collision, ground object'),('Inadvertent and/or incorrect operation of controls.','incorrect operation controls'),('Inability to achieve intended aircraft configuration for any flight phase, e.g. gear, doors, flaps, stabilisers, slats.','configuration, gear, flaps, slats'),('Abnormal vibration.','abnormal vibration'),('Operation of primary warning system associated with manoeuvring, e.g. configuration, stall/stick shake, overspeed warning unless conclusively false and not hazardous.','stall warning, stick shaker, overspeed warning, configuration warning'),('GPWS warning.','GPWS, EGPWS'),('ACAS/TCAS resolution advisories.','ACAS RA, TCAS RA'),('Jet or prop blast incidents causing significant damage or serious injury.','jet blast, prop blast')]
for c,k in ops: add('Aircraft Flight Operations','II-A','Operation of Aircraft',c,k)
for c,k in [('Fire, explosion, smoke or toxic/noxious fumes.','fire, explosion, smoke, fumes'),('Use of non-standard procedure by flight/cabin crew to deal with emergency.','non-standard procedure, emergency'),('Event leading to emergency evacuation.','emergency evacuation'),('Depressurisation.','depressurization'),('Use of emergency equipment or prescribed emergency procedures to deal with a situation.','emergency equipment, emergency procedure'),('Event leading to declaration of emergency.','declared emergency'),('Failure of emergency system/equipment including exit doors.','emergency system failure, exit door'),('Events requiring emergency oxygen use by any crew member.','emergency oxygen')]: add('Aircraft Flight Operations','II-B','Emergencies',c,k)
for c,k in [('Incapacitation of any flight crew member.','flight crew incapacitation'),('Incapacitation of cabin crew member rendering them unable to perform essential emergency duties.','cabin crew incapacitation')]: add('Aircraft Flight Operations','II-C','Crew Incapacitation',c,k)
for c,k in [('Lightning strike causing aircraft damage or loss/malfunction of essential service.','lightning strike'),('Hail strike causing aircraft damage or loss/malfunction of essential service.','hail strike'),('Severe turbulence causing injury to occupants or requiring turbulence check.','severe turbulence, injury'),('Windshear encounter.','windshear'),('Icing encounter causing handling difficulties, aircraft damage or loss/malfunction of essential service.','icing encounter')]: add('Aircraft Flight Operations','II-D','Meteorology',c,k)

# III Maintenance CAW
for c,k in [('Incorrect assembly of aircraft parts/components found during inspection or test.','incorrect assembly'),('Hot bleed air leak resulting in structural damage.','hot bleed air leak'),('Damage/deterioration such as fractures, cracks, corrosion, delamination or disbonding to primary structure/principal structural element requiring repair or replacement.','fracture, crack, corrosion, delamination, primary structure'),('Damage/deterioration to secondary structure which may have endangered aircraft.','secondary structure damage'),('Damage/deterioration to engine, propeller or rotorcraft rotor system.','engine damage, propeller, rotor'),('Products, parts, appliances or materials of unknown/suspect origin.','suspect part, unknown origin'),('Misleading, incorrect or insufficient maintenance data/procedures that could lead to maintenance errors.','maintenance data, procedure error')]: add('Maintenance and Continuing Airworthiness Management','III','Maintenance/Continuing Airworthiness',c,k)

# IV ANS Ground
for c,k in [('Significantly incorrect, inadequate or misleading information from ground sources such as ATC, ATIS, meteorology, navigation databases, maps, charts or manuals.','ATC, ATIS, misleading information'),('Provision of less than prescribed terrain clearance.','terrain clearance'),('Incorrect pressure reference data / altimeter setting.','altimeter setting, QNH'),('Incorrect transmission, receipt or interpretation of significant messages resulting in hazardous situation.','message error, hazardous'),('Separation minima infringement.','separation minima'),('Unauthorised penetration of airspace.','airspace penetration'),('Unlawful radio communication transmission.','unlawful radio'),('Significant degradation/failure of CNS facilities.','CNS failure'),('Aerodrome movement areas obstructed by aircraft, vehicles, animals or foreign objects causing hazardous/potentially hazardous situation.','movement area obstruction, FOD, animals'),('Errors/inadequacies in marking obstructions or hazards on aerodrome movement areas causing hazardous situation.','marking error, obstruction'),('Failure, significant malfunction or unavailability of airfield lighting.','airfield lighting failure')]: add('Air Navigation Services, Facilities and Ground Services','IV-A','Air Navigation Services',c,k)
for c,k in [('Significant spillage during fueling operations.','fuel spillage'),('Loading of incorrect fuel quantities likely to significantly affect endurance, performance, balance or structural strength.','incorrect fuel quantity'),('Unsatisfactory ground de-icing / anti-icing.','ground deicing, anti icing')]: add('Air Navigation Services, Facilities and Ground Services','IV-B','Aerodrome and Aerodrome Facilities',c,k)
for c,k in [('Significant contamination of aircraft structure/systems/equipment from carriage of baggage or cargo.','baggage contamination, cargo contamination'),('Incorrect loading of passengers, baggage or cargo likely to significantly affect mass and/or balance.','incorrect loading, mass balance'),('Incorrect stowage of baggage/cargo including hand baggage likely to create hazardous situation or impede emergency evacuation.','incorrect stowage, evacuation impeded'),('Inadequate stowage of cargo containers or substantial cargo items.','cargo container stowage'),('Dangerous goods incidents.','dangerous goods, DG')]: add('Air Navigation Services, Facilities and Ground Services','IV-C','Passenger Handling, Baggage and Cargo',c,k)
for c,k in [('Failure/malfunction/defect of ground equipment used for aircraft system/equipment testing/checking where routine inspection/test procedures did not identify problem and results in hazardous situation.','ground equipment failure, test equipment'),('Loading of contaminated or incorrect type of fuel or other essential fluids including oxygen and potable water.','contaminated fuel, incorrect fluid, oxygen, potable water')]: add('Air Navigation Services, Facilities and Ground Services','IV-D','Aircraft Ground Handling and Servicing',c,k)

# V-VIII IX
add('Maintenance Organization','V','Maintenance Organization','Any airframe, engine, propeller, component or system defect/malfunction/damage found during scheduled or unscheduled maintenance that could possibly lead to aircraft operational accident or serious incident if not properly rectified.','maintenance defect, malfunction, damage')
add('Design and Manufacturing Organizations','VI','Design/Manufacturing','Design/manufacturing deficiency, defect or malfunction of product/services warranting possible issue of EAD, AD or ASB.','design defect, manufacturing defect, EAD, AD, ASB')
add('Wildlife Activity','VII','Wildlife Activity','All wildlife strikes and wildlife movement are reportable in the prescribed performa.','wildlife strike, bird strike, animal movement')
add('Ground Collision Incidents','VIII','Ground Collision','Collision while taxiing to or from a runway in use.','taxiing collision, runway in use')
add('Ground Handling Ramp Incidents','IX','Ground Handling Ramp','Ground handling/ramp incidents are listed as reportable occurrence heading; specific conditions should be defined by organization based on approved procedures/manuals.','ramp incident, ground handling','PDF heading present but no detailed bullet conditions in extracted text.')

df = pd.DataFrame(rules)
summary = df.groupby(['domain','section','subcategory']).size().reset_index(name='rule_count')

out='rules/structured_MOR_rules.xlsx'
with pd.ExcelWriter(out, engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='MOR_Rules', index=False)
    summary.to_excel(writer, sheet_name='Summary', index=False)
    pd.DataFrame({
        'field':['rule_id','domain','section','subcategory','mor_condition','suggested_label','keywords_or_triggers','notes_or_exceptions'],
        'description':['Unique rule identifier','Top-level CAR Appendix A domain','Document section code','Detailed category/subcategory','Extracted reportable occurrence condition','Default label for ML/rule engine','Practical keyword triggers for matching','Exceptions/implementation notes']
    }).to_excel(writer, sheet_name='Data_Dictionary', index=False)

wb=load_workbook(out)
for ws in wb.worksheets:
    ws.freeze_panes='A2'
    for cell in ws[1]:
        cell.font=Font(bold=True, color='FFFFFF')
        cell.fill=PatternFill('solid', fgColor='1F4E78')
        cell.alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(vertical='top', wrap_text=True)
    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len+2, 12), 55)
wb.save(out)
print(out, len(df), 'rules created')