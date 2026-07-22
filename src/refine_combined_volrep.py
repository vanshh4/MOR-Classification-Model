"""
refine_combined_volrep.py

Purpose:
    Further clean the already processed Volrep workbook by consolidating:
        1. Multiple event date/time columns into one column: event_date_time
        2. Multiple place-of-occurrence/location columns into one column: place_of_occurrence

Input:
    data/processed/volrep/combined_volrep_2026H1_cleaned.xlsx

Output:
    data/processed/volrep/combined_volrep_2026H1_refined.xlsx

Run from project root:
    python src/refine_combined_volrep.py

Notes:
    - This script does not change raw_report_text or clean_report_text.
    - It removes redundant date/time and place/location columns after creating the consolidated columns.
    - It preserves important operational fields such as departure station, airfield of landing, aircraft type,
      phase of flight, and cleaned report text.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# 1. PROJECT PATH CONFIGURATION
# ============================================================


def find_project_root() -> Path:
    """
    Find project root robustly.

    Expected location:
        MOR_Classification_Project/src/refine_combined_volrep.py

    If run as a standalone file, it falls back to current working directory.
    """
    script_path = Path(__file__).resolve()

    if script_path.parent.name.lower() == "src":
        return script_path.parents[1]

    for parent in [script_path.parent] + list(script_path.parents):
        if (parent / "data" / "processed" / "volrep").exists():
            return parent

    return Path.cwd()


PROJECT_ROOT = find_project_root()

PROCESSED_VOLREP_DIR = PROJECT_ROOT / "data" / "processed" / "volrep"

INPUT_FILE = PROCESSED_VOLREP_DIR / "combined_volrep_2026H1_cleaned.xlsx"
FALLBACK_INPUT_FILE = Path("combined_volrep_2026H1_cleaned.xlsx")

OUTPUT_FILE = PROCESSED_VOLREP_DIR / "combined_volrep_2026H1_refined.xlsx"
FALLBACK_OUTPUT_FILE = Path("combined_volrep_2026H1_refined.xlsx")


# ============================================================
# 2. COLUMN CONFIGURATION
# ============================================================

DATE_TIME_COLUMNS = [
    "event_datetime_parsed",
    "event_date",
    "Date and Time of Occurrence (24Hr format and UTC)",
    "Date and Time of Occurrence (24Hr format)",
    "Date & time of event",
    "Date_2",
]

PLACE_COLUMNS_PRIORITY = [
    "place_of_occurrence_normalized",
    "place_of_occurrence_2_normalized",
    "location_normalized",
    "Place of Occurrence",
    "Place of Occurrence_2",
    "Location",
]

PLACE_COLUMNS_TO_DROP = [
    "Place of Occurrence",
    "Place of Occurrence_2",
    "Location",
    "place_of_occurrence_normalized",
    "place_of_occurrence_code",
    "place_of_occurrence_2_normalized",
    "place_of_occurrence_2_code",
    "location_normalized",
    "location_code",
]

# These are not place-of-occurrence columns. They are route/station context, so they are preserved.
ROUTE_STATION_COLUMNS_TO_KEEP = [
    "Departure Station",
    "Airfield of Landing",
    "Arrival station",
    "departure_station_normalized",
    "departure_station_code",
    "airfield_of_landing_normalized",
    "airfield_of_landing_code",
    "arrival_station_normalized",
    "arrival_station_code",
]


# ============================================================
# 3. HELPER FUNCTIONS
# ============================================================


def resolve_input_path() -> Path:
    """Resolve input path either from project hierarchy or local fallback."""
    if INPUT_FILE.exists():
        return INPUT_FILE
    if FALLBACK_INPUT_FILE.exists():
        return FALLBACK_INPUT_FILE
    raise FileNotFoundError(
        "Input file not found. Checked:\n"
        f"1. {INPUT_FILE}\n"
        f"2. {FALLBACK_INPUT_FILE}"
    )


def resolve_output_path() -> Path:
    """Resolve output path. Prefer project hierarchy, otherwise local fallback."""
    try:
        PROCESSED_VOLREP_DIR.mkdir(parents=True, exist_ok=True)
        return OUTPUT_FILE
    except PermissionError:
        return FALLBACK_OUTPUT_FILE


def is_blank_or_unknown(value: object) -> bool:
    """Check if value should be treated as missing/unknown."""
    if pd.isna(value):
        return True
    text = str(value).strip()
    return text == "" or text.upper() in {"UNKNOWN", "NAN", "NA", "N/A", "NONE", "NULL", "-"}


def clean_station_text(value: object) -> str:
    """
    Normalize station/place text lightly.

    Example:
        VABB : MUMBAI,CHHATRAPATI SHIVAJI INTL
        -> VABB - MUMBAI CHHATRAPATI SHIVAJI INTL
    """
    if is_blank_or_unknown(value):
        return "UNKNOWN"

    text = str(value).replace("\u00a0", " ").strip().upper()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s*,\s*", " ", text)
    text = re.sub(r"\s*:\s*", " : ", text)
    text = re.sub(r"\s+", " ", text).strip()

    match = re.match(r"^([A-Z]{4})\s*:\s*(.+)$", text)
    if match:
        code = match.group(1).strip()
        name = match.group(2).strip()
        return f"{code} - {name}"

    return text


def first_valid_value(row: pd.Series, columns: Iterable[str]) -> object:
    """Return first non-empty/non-UNKNOWN value available in the row."""
    for col in columns:
        if col in row.index and not is_blank_or_unknown(row[col]):
            return row[col]
    return np.nan


def parse_possible_excel_serial(value: object) -> pd.Timestamp:
    """
    Parse dates that may appear as:
        - pandas Timestamp
        - Excel serial number, e.g. 46022
        - string date, e.g. 31-Dec-25 or 2026/03/30 12:55:00
    """
    if is_blank_or_unknown(value):
        return pd.NaT

    if isinstance(value, pd.Timestamp):
        return value

    # Excel serial dates generally appear as integers/floats around 40000-50000.
    # Excel origin 1899-12-30 is used by pandas for typical Excel serial conversion.
    if isinstance(value, (int, float, np.integer, np.floating)):
        try:
            numeric_value = float(value)
            if 20000 <= numeric_value <= 70000:
                return pd.to_datetime(numeric_value, unit="D", origin="1899-12-30", errors="coerce")
        except Exception:
            pass

    text = str(value).strip()

    # Sometimes Excel serial is read as string.
    if re.fullmatch(r"\d+(\.0)?", text):
        try:
            numeric_value = float(text)
            if 20000 <= numeric_value <= 70000:
                return pd.to_datetime(numeric_value, unit="D", origin="1899-12-30", errors="coerce")
        except Exception:
            pass

    # Generic parsing fallback.
    return pd.to_datetime(text, errors="coerce", dayfirst=False)


def coalesce_event_datetime(row: pd.Series) -> pd.Timestamp:
    """Create a single event date-time value from multiple possible date columns."""
    for col in DATE_TIME_COLUMNS:
        if col in row.index:
            parsed = parse_possible_excel_serial(row[col])
            if pd.notna(parsed):
                return parsed
    return pd.NaT


def coalesce_place_of_occurrence(row: pd.Series) -> str:
    """Create a single place_of_occurrence value from multiple possible location columns."""
    value = first_valid_value(row, PLACE_COLUMNS_PRIORITY)
    return clean_station_text(value)


def format_event_datetime(value: pd.Timestamp) -> str:
    """Format event date-time as YYYY-MM-DD HH:MM:SS. If time is absent, use 00:00:00."""
    if pd.isna(value):
        return ""
    return pd.Timestamp(value).strftime("%Y-%m-%d %H:%M:%S")


# ============================================================
# 4. MAIN REFINING LOGIC
# ============================================================


def refine_volrep_file() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Load, refine and return cleaned dataframe plus quality metadata."""
    input_path = resolve_input_path()

    df = pd.read_excel(input_path, sheet_name="Cleaned_Volrep", engine="openpyxl")

    original_columns = list(df.columns)
    original_row_count = len(df)

    # Create one consolidated event date-time column.
    df["event_date_time"] = df.apply(coalesce_event_datetime, axis=1)
    df["event_date_time"] = df["event_date_time"].apply(format_event_datetime)

    # Create one consolidated place of occurrence column.
    df["place_of_occurrence"] = df.apply(coalesce_place_of_occurrence, axis=1)

    # Count quality indicators before dropping columns.
    missing_event_datetime = int(df["event_date_time"].eq("").sum())
    missing_place = int(df["place_of_occurrence"].eq("UNKNOWN").sum())

    # Drop redundant date/time and place columns.
    columns_to_drop = [
        col for col in (DATE_TIME_COLUMNS + PLACE_COLUMNS_TO_DROP)
        if col in df.columns
    ]

    df = df.drop(columns=columns_to_drop, errors="ignore")

    # Reorder important columns toward the front.
    preferred_columns = [
        "volrep_internal_id",
        "Number",
        "source_file",
        "source_period",
        "Reportee Role",
        "aircraft_type_normalized",
        "phase_of_flight_normalized",
        "event_date_time",
        "place_of_occurrence",
        "raw_report_text",
        "clean_report_text",
        "raw_text_word_count",
        "clean_text_word_count",
        "is_text_missing",
        "is_short_text",
        "duplicate_clean_text",
    ]

    preferred_existing = [col for col in preferred_columns if col in df.columns]
    remaining_columns = [col for col in df.columns if col not in preferred_existing]
    df = df[preferred_existing + remaining_columns]

    removed_columns_df = pd.DataFrame({"removed_column": columns_to_drop})

    quality_summary = pd.DataFrame([
        {"metric": "input_file", "value": str(input_path)},
        {"metric": "output_file", "value": str(resolve_output_path())},
        {"metric": "original_row_count", "value": original_row_count},
        {"metric": "final_row_count", "value": len(df)},
        {"metric": "original_column_count", "value": len(original_columns)},
        {"metric": "final_column_count", "value": len(df.columns)},
        {"metric": "removed_column_count", "value": len(columns_to_drop)},
        {"metric": "missing_event_date_time_rows", "value": missing_event_datetime},
        {"metric": "unknown_place_of_occurrence_rows", "value": missing_place},
    ])

    return df, quality_summary, removed_columns_df


def save_refined_workbook(
    refined_df: pd.DataFrame,
    quality_summary: pd.DataFrame,
    removed_columns_df: pd.DataFrame,
) -> Path:
    """Save refined workbook with supporting summary sheets."""
    output_path = resolve_output_path()

    data_dictionary = pd.DataFrame([
        {
            "column": "event_date_time",
            "description": "Single consolidated event date/time column created from all available date/time columns.",
        },
        {
            "column": "place_of_occurrence",
            "description": "Single consolidated place/location column created from available place/location fields.",
        },
        {
            "column": "clean_report_text",
            "description": "ML-ready cleaned event narrative retained from previous preprocessing stage.",
        },
        {
            "column": "raw_report_text",
            "description": "Original coalesced event narrative retained from previous preprocessing stage.",
        },
    ])

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        refined_df.to_excel(writer, sheet_name="Cleaned_Volrep_Refined", index=False)
        quality_summary.to_excel(writer, sheet_name="Quality_Summary", index=False)
        removed_columns_df.to_excel(writer, sheet_name="Removed_Columns", index=False)
        data_dictionary.to_excel(writer, sheet_name="Data_Dictionary", index=False)

    # Light formatting for readability.
    wb = load_workbook(output_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 55)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(output_path)
    return output_path


def main() -> None:
    refined_df, quality_summary, removed_columns_df = refine_volrep_file()
    output_path = save_refined_workbook(refined_df, quality_summary, removed_columns_df)

    print("Combined Volrep refinement completed successfully.")
    print(f"Rows processed: {len(refined_df)}")
    print(f"Output saved to: {output_path}")
    print("\nQuality summary:")
    print(quality_summary.to_string(index=False))


if __name__ == "__main__":
    main()
