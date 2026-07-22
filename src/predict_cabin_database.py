"""
predict_cabin_database.py

Purpose:
    Run the trained MOR binary classifier on every Cabin Database report and
    add a new column named "Prediction results" containing either:
        - MOR
        - Non-MOR

Input:
    data/raw/cabin/CABIN DATABASE 2026.xlsx

Output:
    outputs/predictions/cabin/cabin_database_mor_predictions.xlsx

Run from the project root:
    python src/predict_cabin_database.py

Important:
    - The original workbook is never overwritten.
    - Existing worksheets and cell formatting are preserved using openpyxl.
    - The script accepts common narrative-column variants, including the
      workbook's current "BREIF DESCRIPTION" spelling.
    - Predictions are model-generated recommendations and require human review.
"""

from __future__ import annotations

import inspect
import re
import sys
from pathlib import Path
from typing import Any, Optional

import joblib
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# =============================================================================
# 1. PROJECT PATHS
# =============================================================================


def find_project_root() -> Path:
    """Find the project root when this script is stored in src/."""
    script_path = Path(__file__).resolve()

    if script_path.parent.name.lower() == "src":
        return script_path.parents[1]

    for parent in [script_path.parent, *script_path.parents]:
        if (parent / "src").exists() and (parent / "models").exists():
            return parent

    return Path.cwd()


PROJECT_ROOT = find_project_root()
SRC_DIR = PROJECT_ROOT / "src"

INPUT_FILE = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "cabin"
    / "CABIN DATABASE 2026.xlsx"
)

OUTPUT_DIR = PROJECT_ROOT / "outputs" / "predictions" / "cabin"
OUTPUT_FILE = OUTPUT_DIR / "cabin_database_mor_predictions.xlsx"

BINARY_MODEL_FILE = PROJECT_ROOT / "models" / "mor_binary_model.pkl"
BINARY_VECTORIZER_FILE = PROJECT_ROOT / "models" / "mor_binary_vectorizer.pkl"

# Fine-tuned MOR decision threshold used by the project.
MOR_DECISION_THRESHOLD = 0.45

# Final output column requested by the project.
PREDICTION_COLUMN = "Prediction results"

# The uploaded workbook currently uses "BREIF DESCRIPTION". Other aliases are
# supported to keep the script robust if the heading is corrected later.
TEXT_COLUMN_ALIASES = (
    "Brief Description",
    "BREIF DESCRIPTION",
    "BRIEF DESCRIPTION",
    "brief_description",
    "clean_brief_description",
)


# =============================================================================
# 2. HELPERS
# =============================================================================


def normalize_header(value: Any) -> str:
    """Normalize an Excel heading for reliable alias matching."""
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def normalize_report_text(value: Any) -> str:
    """Apply light whitespace normalization without removing aviation terms."""
    if value is None:
        return ""
    return " ".join(str(value).replace("\u00a0", " ").split()).strip()


def resolve_text_column(headers: list[Any]) -> Optional[int]:
    """Return the 1-based Excel column index for the report narrative."""
    normalized_headers = {
        normalize_header(header): index
        for index, header in enumerate(headers, start=1)
        if header is not None
    }

    for alias in TEXT_COLUMN_ALIASES:
        index = normalized_headers.get(normalize_header(alias))
        if index is not None:
            return index

    return None


def load_binary_artifacts():
    """Load the trained MOR binary model and TF-IDF vectorizer."""
    if not BINARY_MODEL_FILE.exists():
        raise FileNotFoundError(
            f"Binary MOR model not found: {BINARY_MODEL_FILE}\n"
            "Run python src/train_mor_binary.py first."
        )

    if not BINARY_VECTORIZER_FILE.exists():
        raise FileNotFoundError(
            f"Binary MOR vectorizer not found: {BINARY_VECTORIZER_FILE}\n"
            "Run python src/train_mor_binary.py first."
        )

    model = joblib.load(BINARY_MODEL_FILE)
    vectorizer = joblib.load(BINARY_VECTORIZER_FILE)
    return model, vectorizer


def predict_label(report_text: str, model, vectorizer) -> str:
    """Predict MOR or Non-MOR for one normalized report narrative."""
    transformed = vectorizer.transform([report_text])
    probabilities = model.predict_proba(transformed)[0]

    class_to_index = {
        class_label: index
        for index, class_label in enumerate(model.classes_)
    }

    if 1 not in class_to_index:
        raise ValueError(
            f"The trained model does not contain MOR class 1. "
            f"Available classes: {model.classes_.tolist()}"
        )

    mor_probability = float(probabilities[class_to_index[1]])
    return "MOR" if mor_probability >= MOR_DECISION_THRESHOLD else "Non-MOR"


def find_or_create_prediction_column(worksheet) -> int:
    """Find the requested output column or append it to the worksheet."""
    for cell in worksheet[1]:
        if normalize_header(cell.value) == normalize_header(PREDICTION_COLUMN):
            return cell.column

    prediction_column = worksheet.max_column + 1
    header_cell = worksheet.cell(row=1, column=prediction_column)
    header_cell.value = PREDICTION_COLUMN

    # Match the general professional style while making the generated column
    # easy for human reviewers to identify.
    header_cell.fill = PatternFill("solid", fgColor="7030A0")
    header_cell.font = Font(color="FFFFFF", bold=True)
    header_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    worksheet.column_dimensions[get_column_letter(prediction_column)].width = 20

    return prediction_column


# =============================================================================
# 3. WORKBOOK PREDICTION PIPELINE
# =============================================================================


def process_workbook() -> None:
    """Add MOR/Non-MOR predictions to all sheets containing report text."""
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Cabin Database input file not found: {INPUT_FILE}\n"
            "Place the workbook at data/raw/cabin/CABIN DATABASE 2026.xlsx"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    model, vectorizer = load_binary_artifacts()
    workbook = load_workbook(INPUT_FILE)

    total_reports = 0
    total_mor = 0
    total_non_mor = 0
    skipped_blank_text = 0
    processed_sheets: list[str] = []

    for worksheet in workbook.worksheets:
        headers = [cell.value for cell in worksheet[1]]
        text_column = resolve_text_column(headers)

        # Supporting sheets such as Terminology and ERC Score Table do not have
        # report narratives and must remain unchanged.
        if text_column is None:
            continue

        prediction_column = find_or_create_prediction_column(worksheet)
        processed_sheets.append(worksheet.title)

        for row_number in range(2, worksheet.max_row + 1):
            report_text = normalize_report_text(
                worksheet.cell(row=row_number, column=text_column).value
            )

            prediction_cell = worksheet.cell(
                row=row_number,
                column=prediction_column,
            )

            if not report_text:
                # Keep non-report/empty rows blank rather than assigning a label.
                prediction_cell.value = None
                skipped_blank_text += 1
                continue

            prediction = predict_label(report_text, model, vectorizer)
            prediction_cell.value = prediction
            prediction_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            if prediction == "MOR":
                prediction_cell.fill = PatternFill("solid", fgColor="F4CCCC")
                prediction_cell.font = Font(color="9C0006", bold=True)
                total_mor += 1
            else:
                prediction_cell.fill = PatternFill("solid", fgColor="D9EAD3")
                prediction_cell.font = Font(color="274E13", bold=True)
                total_non_mor += 1

            total_reports += 1

        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.freeze_panes = "A2"

    if not processed_sheets:
        raise ValueError(
            "No worksheet contains a supported Brief Description column. "
            f"Supported headings: {', '.join(TEXT_COLUMN_ALIASES)}"
        )

    workbook.save(OUTPUT_FILE)

    print("Cabin Database MOR prediction completed successfully.")
    print(f"Input file: {INPUT_FILE}")
    print(f"Output file: {OUTPUT_FILE}")
    print(f"Processed sheets: {', '.join(processed_sheets)}")
    print(f"Reports processed: {total_reports}")
    print(f"Predicted MOR: {total_mor}")
    print(f"Predicted Non-MOR: {total_non_mor}")
    print(f"Rows skipped because report text was blank: {skipped_blank_text}")
    print("Important: Prediction results require human verification.")


if __name__ == "__main__":
    process_workbook()
